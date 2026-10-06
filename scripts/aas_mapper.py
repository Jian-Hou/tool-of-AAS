"""Generic Excel-to-AAS mapping: read any workbook, open or create an AAS, map cells and rows into it."""
import copy
import datetime
import hashlib
import io
import json
import os
import posixpath
import re
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import unquote, urlparse
from xml.etree import ElementTree as ET

import openpyxl
from aas_core3 import jsonization, types as aas_types, verification, xmlization

from conversion import DECIMAL, MAX_ROWS, MAX_ZIP_BYTES, ValidationError, integer, number, safe_id, text

MAX_EXCEL_BYTES = 12 * 1024 * 1024
MAX_AASX_BYTES = 256 * 1024 * 1024
ID_SHORT = re.compile(r'[A-Za-z][A-Za-z0-9_]{0,127}')
LANGUAGE = re.compile(r'[A-Za-z]{2,3}(-[A-Za-z0-9]{1,8})*')
CHILDREN = {'Submodel': 'submodelElements', 'SubmodelElementCollection': 'value', 'SubmodelElementList': 'value',
            'Entity': 'statements', 'AnnotatedRelationshipElement': 'annotations'}
CONTAINERS = ('Submodel', 'SubmodelElementCollection', 'SubmodelElementList', 'Entity')
VALUE_ELEMENTS = ('Property', 'MultiLanguageProperty')
ROW_ITEMS = ('SubmodelElementCollection', 'Entity')
GLOBAL_ASSET = '#globalAssetId'  # column target that sets the global asset ID of an entity row
NEW_ELEMENTS = ('Property', 'MultiLanguageProperty', 'SubmodelElementCollection', 'SubmodelElementList')
VALUE_TYPES = tuple(member.value for member in aas_types.DataTypeDefXSD)
RELATIONSHIPS = 'http://schemas.openxmlformats.org/package/2006/relationships'


class Uncached(str):
    """Text of a formula cell for which the workbook stores no calculated result."""


# ---------------------------------------------------------------- Excel tables
@dataclass
class Tables:
    sheets: dict
    source_name: str
    source_sha256: str


def guess_type(values):
    values = [v for v in values if v not in (None, '') and not isinstance(v, Uncached)]
    if not values:
        return 'xs:string'
    if all(isinstance(v, bool) for v in values):
        return 'xs:boolean'
    if all(isinstance(v, int) and not isinstance(v, bool) for v in values):
        return 'xs:integer'
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
        return 'xs:double'
    if all(isinstance(v, datetime.datetime) for v in values):
        return 'xs:date' if all(v.time() == datetime.time() for v in values) else 'xs:dateTime'
    if all(isinstance(v, str) and DECIMAL.fullmatch(v.strip()) for v in values):
        return 'xs:double'
    return 'xs:string'


def read_tables(path, source_name=None):
    """Read every worksheet whose first row holds column headers; problems are reported per sheet."""
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValidationError('Only .xlsx is supported. Save legacy .xls files as .xlsx first.')
    payload = path.read_bytes()
    if len(payload) > MAX_EXCEL_BYTES:
        raise ValidationError('The Excel file exceeds the 12 MB limit.')
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_ZIP_BYTES or len(archive.infolist()) > 1000:
                raise ValidationError('The Excel archive exceeds the supported limits (64 MB / 1,000 entries).')
        formulas = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        values = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Unable to read Excel. Check that the file is a valid .xlsx workbook.') from exc
    sheets = {}
    try:
        if len(formulas.sheetnames) > 32:
            raise ValidationError('The workbook contains more than 32 worksheets.')
        for ws_formula, ws_value in zip(formulas, values):
            sheet = dict(columns=[], rows=[], problem=None, notes=[])
            sheets[ws_formula.title] = sheet
            rows_formula, rows_value = ws_formula.iter_rows(), ws_value.iter_rows()
            names = [text(c.value) for c in next(rows_formula, ())]
            next(rows_value, None)
            while names and not names[-1]:
                names.pop()
            if not names:
                sheet['problem'] = 'The first row contains no column headers.'
                continue
            if any(not n for n in names) or len(set(names)) != len(names):
                sheet['problem'] = 'The header row contains blank or duplicate column names.'
                continue
            sheet['columns'] = names
            cached = uncached = ignored = 0
            for line, (cells, cell_values) in enumerate(zip(rows_formula, rows_value), 2):
                if line > MAX_ROWS or len(cells) > 128:
                    raise ValidationError(f'{ws_formula.title} exceeds the supported limits (25,000 rows / 128 columns).')
                row = {}
                for i, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type == 'f':
                        value = cell_values[i].value if i < len(cell_values) else None
                        if value is None:
                            value, uncached = Uncached(cell.value), uncached + 1
                        else:
                            cached += 1
                    if i >= len(names):
                        ignored += value not in (None, '')
                    elif value not in (None, ''):
                        row[names[i]] = value
                if row:
                    row['_row'] = line
                    sheet['rows'].append(row)
            if cached:
                sheet['notes'].append(f'{cached} formula cells use the result last saved by Excel.')
            if uncached:
                sheet['notes'].append(f'{uncached} formula cells have no saved result and cannot be imported; open and save the file in Excel or paste values.')
            if ignored:
                sheet['notes'].append(f'{ignored} cells in columns without a header are ignored.')
    finally:
        formulas.close()
        values.close()
    return Tables(sheets, source_name or path.name, hashlib.sha256(payload).hexdigest())


def table_summary(tables, preview_rows=8):
    result = {}
    for name, sheet in tables.sheets.items():
        result[name] = dict(columns=sheet['columns'], problem=sheet['problem'], notes=sheet['notes'], row_count=len(sheet['rows']),
                            types={c: guess_type([r.get(c) for r in sheet['rows']]) for c in sheet['columns']},
                            preview=[{k: display_cell(v) for k, v in row.items()} for row in sheet['rows'][:preview_rows]])
    return result


def display_cell(value):
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return value if isinstance(value, (int, float)) or value is None else str(value)


# ---------------------------------------------------------------- AAS packages
@dataclass
class Package:
    env: dict
    spec_part: str      # ZIP entry of the AAS model part; '' when the source was a plain JSON file
    spec_format: str    # 'json' or 'xml'
    filename: str


def canonical(env):
    """Round-trip through aas-core3 so that every stored environment has one normalized JSON form."""
    try:
        return jsonization.to_jsonable(jsonization.environment_from_jsonable(prune(copy.deepcopy(env))))
    except Exception as exc:
        raise ValidationError('The AAS structure is invalid: ' + str(exc)[:300]) from exc


def prune(node):
    # The metamodel forbids empty lists; they appear after deleting the last child of an element.
    if isinstance(node, dict):
        for key in [k for k, v in node.items() if isinstance(v, list) and not v]:
            del node[key]
        for value in node.values():
            prune(value)
    elif isinstance(node, list):
        for value in node:
            prune(value)
    return node


def parse_spec(name, data):
    try:
        if name.lower().endswith('.json'):
            env = jsonization.environment_from_jsonable(json.loads(data.decode('utf-8-sig')))
            fmt = 'json'
        else:
            env = xmlization.environment_from_str(data.decode('utf-8-sig'))
            fmt = 'xml'
    except Exception as exc:
        if re.search(rb'admin-shell\.io/aas/[12]/', data[:4096]):
            raise ValidationError('This is an AAS V1/V2 model. Open it in AASX Package Explorer, save it as AAS V3.0 and load it again.') from exc
        if re.search(rb'admin-shell\.io/aas/3/[1-9]', data[:4096]):
            raise ValidationError('This model uses AAS metamodel V3.1 or later; the tool supports V3.0. For IDTA templates, use the file without "forAASMetamodelV3.1".') from exc
        raise ValidationError('Unable to read the AAS model. Only AAS V3.0 models in XML or JSON are supported. Details: ' + str(exc)[:300]) from exc
    return jsonization.to_jsonable(env), fmt


def serialize(env, fmt):
    model = jsonization.environment_from_jsonable(prune(copy.deepcopy(env)))
    if fmt == 'xml':
        return ('<?xml version="1.0" encoding="utf-8"?>\n' + xmlization.to_str(model)).encode('utf-8')
    return json.dumps(jsonization.to_jsonable(model), ensure_ascii=False, indent=2).encode('utf-8')


def _relationship_targets(archive, source, suffix):
    rels = posixpath.join(posixpath.dirname(source), '_rels', posixpath.basename(source) + '.rels') if source else '_rels/.rels'
    if rels not in archive.namelist():
        return []
    try:
        root = ET.fromstring(archive.read(rels))
    except ET.ParseError as exc:
        raise ValidationError(f'The package relationship file {rels} is invalid.') from exc
    targets = []
    for rel in root.iter(f'{{{RELATIONSHIPS}}}Relationship'):
        if rel.get('Type', '').endswith(suffix) and rel.get('TargetMode') != 'External':
            target = unquote(rel.get('Target', ''))
            targets.append(target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(source), target)))
    return targets


def load_source(path, filename=None):
    """Open an AASX package or a plain AAS JSON environment file."""
    path = Path(path)
    filename = filename or path.name
    if path.stat().st_size > MAX_AASX_BYTES:
        raise ValidationError('The AAS file exceeds the 256 MB limit.')
    if path.suffix.lower() == '.json':
        env, fmt = parse_spec('model.json', path.read_bytes())
        return Package(env, '', fmt, filename)
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValidationError('The file is not a valid AASX package.') from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > 20000 or sum(i.file_size for i in infos) > 4 * MAX_AASX_BYTES:
            raise ValidationError('The AASX package exceeds the supported limits.')
        origins = _relationship_targets(archive, '', 'relationships/aasx-origin')
        if len(origins) != 1:
            raise ValidationError('The package has no AASX origin part.')
        specs = _relationship_targets(archive, origins[0], 'relationships/aas-spec')
        if len(specs) != 1:
            raise ValidationError(f'The package must reference exactly one AAS model part; it references {len(specs)}.')
        if specs[0] not in archive.namelist():
            raise ValidationError(f'The AAS model part {specs[0]} referenced by the package is missing.')
        if archive.getinfo(specs[0]).file_size > MAX_AASX_BYTES:
            raise ValidationError('The AAS model part exceeds the 256 MB limit.')
        env, fmt = parse_spec(specs[0], archive.read(specs[0]))
    return Package(env, specs[0], fmt, filename)


def write_package(path, env, base=None, package=None):
    """Write env as an AASX. With a base package, only the AAS model part is replaced; all other parts are copied unchanged."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    patch = base is not None and package is not None and package.spec_part
    fmt = package.spec_format if patch else 'json'
    body = serialize(env, fmt)
    fd, tmp = tempfile.mkstemp(prefix='.aas_', suffix='.aasx', dir=path.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as out:
            if patch:
                with zipfile.ZipFile(base) as source:
                    for info in source.infolist():
                        if info.filename == package.spec_part:
                            replacement = zipfile.ZipInfo(info.filename, date_time=datetime.datetime.now().timetuple()[:6])
                            replacement.compress_type = zipfile.ZIP_DEFLATED
                            out.writestr(replacement, body)
                        else:
                            out.writestr(info, source.read(info.filename))
            else:
                _write_new_package(out, body)
        spec = package.spec_part if patch else 'aasx/data.json'
        with zipfile.ZipFile(tmp) as check:
            if check.testzip() is not None or parse_spec(spec, check.read(spec))[0] != canonical(env):
                raise ValidationError('The exported file failed read-back verification.')
        os.replace(tmp, path)
    finally:
        Path(tmp).unlink(missing_ok=True)


def _write_new_package(out, body):
    def rels(items):
        root = ET.Element('Relationships', xmlns=RELATIONSHIPS)
        for i, (target, kind) in enumerate(items, 1):
            ET.SubElement(root, 'Relationship', Id=f'R{i}', Type=kind, Target=target)
        return ET.tostring(root, encoding='utf-8', xml_declaration=True)
    types = ET.Element('Types', xmlns='http://schemas.openxmlformats.org/package/2006/content-types')
    for extension, content_type in [('rels', 'application/vnd.openxmlformats-package.relationships+xml'), ('json', 'application/json')]:
        ET.SubElement(types, 'Default', Extension=extension, ContentType=content_type)
    ET.SubElement(types, 'Override', PartName='/aasx/aasx-origin', ContentType='text/plain')
    out.writestr('[Content_Types].xml', ET.tostring(types, encoding='utf-8', xml_declaration=True))
    out.writestr('_rels/.rels', rels([('/aasx/aasx-origin', 'http://admin-shell.io/aasx/relationships/aasx-origin')]))
    out.writestr('aasx/aasx-origin', '')
    out.writestr('aasx/_rels/aasx-origin.rels', rels([('/aasx/data.json', 'http://admin-shell.io/aasx/relationships/aas-spec')]))
    out.writestr('aasx/data.json', body)


def verify(env):
    try:
        model = jsonization.environment_from_jsonable(prune(copy.deepcopy(env)))
    except Exception as exc:
        return ['Unable to parse the AAS structure: ' + str(exc)[:300]]
    return [f'{e.path}: {e.cause}' for e in verification.verify(model)]


def new_failures(before, after):
    """Metamodel failures introduced by an edit; problems the source AAS already had are not blamed on the import."""
    known = set(verify(before))
    return [failure for failure in verify(after) if failure not in known]


# ---------------------------------------------------------------- structure
def new_id():
    return 'urn:uuid:' + str(uuid.uuid4())


def check_id_short(value):
    value = text(value)
    if not ID_SHORT.fullmatch(value):
        raise ValidationError(f'Invalid idShort {value!r}: use letters, digits and underscores, start with a letter, at most 128 characters.')
    return value


def external_reference(value):
    value = text(value)
    if len(value) > 2000 or any(ch.isspace() for ch in value):
        raise ValidationError('A semantic ID must be a single identifier without spaces.')
    return {'type': 'ExternalReference', 'keys': [{'type': 'GlobalReference', 'value': value}]}


def new_environment(id_short, global_asset_id=''):
    id_short = check_id_short(id_short)
    global_asset_id = text(global_asset_id)
    if global_asset_id and (not urlparse(global_asset_id).scheme or re.search(r'\s', global_asset_id)):
        raise ValidationError('The global asset ID must be a URI without whitespace, such as urn:... or https://....')
    env = {'assetAdministrationShells': [{'modelType': 'AssetAdministrationShell', 'id': new_id(), 'idShort': id_short,
                                          'assetInformation': {'assetKind': 'Instance', 'globalAssetId': global_asset_id or new_id()}}]}
    return canonical(env)


def shells(env):
    return [{'id': s['id'], 'idShort': s.get('idShort', ''), 'globalAssetId': s.get('assetInformation', {}).get('globalAssetId', '')}
            for s in env.get('assetAdministrationShells', [])]


def get_shell(env, shell_id):
    shell = next((s for s in env.get('assetAdministrationShells', []) if s['id'] == shell_id), None)
    if shell is None:
        raise ValidationError('Select an asset administration shell first.')
    return shell


def submodels_of(env, shell_id):
    by_id = {s['id']: s for s in env.get('submodels', [])}
    ids = [r['keys'][0]['value'] for r in get_shell(env, shell_id).get('submodels', []) if r.get('keys')]
    return [by_id[i] for i in ids if i in by_id]


def items_of(node):
    return node.get(CHILDREN.get(node.get('modelType'), ''), None) or []


def find_child(node, segment):
    items = items_of(node)
    if node.get('modelType') == 'SubmodelElementList':
        return items[segment] if isinstance(segment, int) and 0 <= segment < len(items) else None
    return next((c for c in items if c.get('idShort') == segment), None)


def resolve(env, shell_id, target, required=True):
    submodels = submodels_of(env, shell_id)
    node = next((s for s in submodels if s['id'] == target.get('submodel')), None)
    if node is None and target.get('submodel_idShort'):
        node = next((s for s in submodels if s.get('idShort') == target['submodel_idShort']), None)
    for segment in target.get('path', []) if node is not None else []:
        node = find_child(node, segment)
        if node is None:
            break
    if node is None and required:
        raise ValidationError('Target not found in the selected AAS: ' + describe_target(target))
    return node


def describe_target(target):
    return '/'.join([str(target.get('submodel_idShort') or target.get('submodel'))] + [f'[{s}]' if isinstance(s, int) else s for s in target.get('path', [])])


def display_value(node):
    kind = node.get('modelType')
    if kind == 'MultiLanguageProperty':
        texts = node.get('value') or []
        value = next((t['text'] for t in texts if t.get('language') == 'en'), texts[0]['text'] if texts else None)
    elif kind == 'Range':
        value = f"{node.get('min', '')} .. {node.get('max', '')}" if 'min' in node or 'max' in node else None
    elif kind == 'ReferenceElement':
        value = ' / '.join(k['value'] for k in (node.get('value') or {}).get('keys', [])) or None
    elif kind == 'Entity':
        value = node.get('globalAssetId')
    elif kind in CONTAINERS:
        value = None
    else:
        value = node.get('value')
    return None if value is None else str(value)[:160]


def cardinality(node):
    return next((q.get('value') for q in node.get('qualifiers', []) if q.get('type') in ('SMT/Cardinality', 'Cardinality', 'Multiplicity')), None)


def tree(env, shell_id, limit=25):
    """Flatten the selected shell's submodels for the page. Large containers show their first children and a summary line."""
    nodes = []
    def visit(submodel, node, path, depth):
        kind = node.get('modelType')
        semantic = (node.get('semanticId') or {}).get('keys', [{}])
        nodes.append(dict(target={'submodel': submodel['id'], 'submodel_idShort': submodel.get('idShort', ''), 'path': path},
                          label=node.get('idShort') or (f'[{path[-1]}]' if path else submodel['id']), modelType=kind,
                          valueType=node.get('valueType'), value=display_value(node), semanticId=semantic[0].get('value') if semantic else None,
                          cardinality=cardinality(node), depth=depth, container=kind in CONTAINERS,
                          listType=node.get('typeValueListElement'), template=submodel.get('kind') == 'Template'))
        children = items_of(node)
        for i, child in enumerate(children[:limit]):
            visit(submodel, child, path + [i if kind == 'SubmodelElementList' else child.get('idShort')], depth + 1)
        if len(children) > limit:
            hidden = len(children) - limit
            nodes.append(dict(target={'submodel': submodel['id'], 'submodel_idShort': submodel.get('idShort', ''), 'path': path + ['…']},
                              label=f'… {hidden} more', modelType='More', valueType=None, value=None, semanticId=None, cardinality=None,
                              depth=depth + 1, container=False, listType=None, template=False, more=hidden))
    for submodel in submodels_of(env, shell_id):
        visit(submodel, submodel, [], 0)
    return nodes


def add_submodel(env, shell_id, id_short, semantic_id=''):
    id_short = check_id_short(id_short)
    if any(s.get('idShort') == id_short for s in submodels_of(env, shell_id)):
        raise ValidationError(f'The AAS already has a submodel named {id_short}.')
    submodel = {'modelType': 'Submodel', 'id': new_id(), 'idShort': id_short, 'kind': 'Instance'}
    if text(semantic_id):
        submodel['semanticId'] = external_reference(semantic_id)
    _attach(env, shell_id, submodel)
    return submodel


def _attach(env, shell_id, submodel):
    env.setdefault('submodels', []).append(submodel)
    get_shell(env, shell_id).setdefault('submodels', []).append({'type': 'ModelReference', 'keys': [{'type': 'Submodel', 'value': submodel['id']}]})


def clear_values(node):
    # Template values are examples or placeholders; imported data must not be mixed with them.
    if isinstance(node, dict):
        if node.get('modelType') in ('Property', 'MultiLanguageProperty', 'File', 'Blob', 'ReferenceElement'):
            node.pop('value', None)
            node.pop('valueId', None)
        if node.get('modelType') == 'Range':
            node.pop('min', None)
            node.pop('max', None)
        for value in node.values():
            clear_values(value)
    elif isinstance(node, list):
        for value in node:
            clear_values(value)
    return node


PLACEHOLDER_ASSET = 'https://admin-shell.io/'
LEAVES = ('Property', 'MultiLanguageProperty', 'File', 'Blob', 'ReferenceElement', 'Range')


def template_cardinality(node):
    """Cardinality from a template qualifier; only elements that came from a template carry one."""
    return next((q.get('value') for q in node.get('qualifiers', []) if q.get('kind') == 'TemplateQualifier'
                 and q.get('type') in ('SMT/Cardinality', 'Cardinality', 'Multiplicity')), None)


def has_data(node):
    kind = node.get('modelType')
    if kind == 'Range':
        return 'min' in node or 'max' in node
    if kind in LEAVES:
        return bool(node.get('value') or node.get('valueId'))
    if kind == 'Entity' and node.get('globalAssetId') and not node['globalAssetId'].startswith(PLACEHOLDER_ASSET):
        return True
    if kind in CHILDREN:
        return any(has_data(child) for child in items_of(node))
    # Relationships in a template point to placeholders; other element kinds are kept as they are.
    return kind not in ('RelationshipElement', 'AnnotatedRelationshipElement')


def replaced_placeholder(node, siblings):
    """An empty template element that repeats (OneToMany) and already has filled entries of the same kind,
    such as the template's Node entity after rows were imported with it as the row structure."""
    def same_kind(other):
        return (other is not node and other.get('modelType') == node.get('modelType') and other.get('semanticId') == node.get('semanticId')
                and template_cardinality(other) == template_cardinality(node) and has_data(other))
    return template_cardinality(node) == 'OneToMany' and not has_data(node) and any(same_kind(other) for other in siblings)


def _drop_empty_optional(node, stats):
    key = CHILDREN.get(node.get('modelType'))
    if not key or not node.get(key):
        return
    kept = []
    for child in node[key]:
        if (template_cardinality(child) in ('ZeroToOne', 'ZeroToMany') and not has_data(child)) or replaced_placeholder(child, node[key]):
            stats['removed'] += 1
            continue
        _drop_empty_optional(child, stats)
        kept.append(child)
    node[key] = kept


def _requirement_warnings(submodel, found):
    """Collect (message pattern, path) pairs; list indices become [*] so repeated entries can be counted together."""
    def visit(node, path):
        kind = node.get('modelType')
        cardinality = template_cardinality(node)
        where = '/'.join([submodel.get('idShort', submodel['id'])] + ['[*]' if isinstance(s, int) else s for s in path])
        listed = any(isinstance(s, int) for s in path)
        if path and not has_data(node):
            if cardinality in ('One', 'OneToMany'):
                # One warning per empty mandatory element; its own children are covered by it.
                found.append((f'Mandatory template element {where} ({cardinality}) ' + ('has no value' if kind in LEAVES else 'is empty'), listed))
                return
            if cardinality in ('ZeroToOne', 'ZeroToMany'):
                return  # an unused optional group does not make its mandatory children required
        if path and kind == 'Entity' and (node.get('globalAssetId') or '').startswith(PLACEHOLDER_ASSET):
            found.append((f'Entity {where} still uses the template placeholder asset ID {node["globalAssetId"]}', listed))
        for i, child in enumerate(items_of(node)):
            if not replaced_placeholder(child, items_of(node)):
                visit(child, path + [i if kind == 'SubmodelElementList' else child.get('idShort')])
    visit(submodel, [])


def finalize(env, drop_empty_optional=True):
    """Prepare a working environment for export: optionally remove empty optional template elements,
    report mandatory template elements without data, and remove template-only qualifiers (AASd-129)."""
    final = copy.deepcopy(env)
    stats = {'removed': 0}
    if drop_empty_optional:
        for submodel in final.get('submodels', []):
            _drop_empty_optional(submodel, stats)
    found = []
    for submodel in final.get('submodels', []):
        _requirement_warnings(submodel, found)
    counts = {}
    for message, listed in found:
        counts[message] = (counts.get(message, (0, listed))[0] + 1, listed)
    warnings = [message + (f' in {count} list entries.' if listed else '.') for message, (count, listed) in counts.items()]
    strip_template_qualifiers(final)
    return prune(final), {'removed': stats['removed'], 'warnings': warnings}


def strip_template_qualifiers(node):
    # Template qualifiers such as SMT/Cardinality are only allowed inside submodel templates (AASd-129).
    if isinstance(node, dict):
        if 'qualifiers' in node:
            node['qualifiers'] = [q for q in node['qualifiers'] if q.get('kind') != 'TemplateQualifier']
        for value in node.values():
            strip_template_qualifiers(value)
    elif isinstance(node, list):
        for value in node:
            strip_template_qualifiers(value)
    return node


def _replace_reference_targets(node, old, new):
    if isinstance(node, dict):
        if node.get('type') == 'ModelReference' and node.get('keys') and node['keys'][0].get('value') == old:
            node['keys'][0]['value'] = new
        for value in node.values():
            _replace_reference_targets(value, old, new)
    elif isinstance(node, list):
        for value in node:
            _replace_reference_targets(value, old, new)


def template_submodels(env):
    return [{'id': s['id'], 'idShort': s.get('idShort', ''), 'kind': s.get('kind', 'Instance'),
             'semanticId': ((s.get('semanticId') or {}).get('keys') or [{}])[0].get('value')} for s in env.get('submodels', [])]


def add_template_submodel(env, shell_id, template_env, submodel_id, keep_values=False):
    source = next((s for s in template_env.get('submodels', []) if s['id'] == submodel_id), None)
    if source is None:
        raise ValidationError('The selected template submodel was not found.')
    if any(s.get('idShort') == source.get('idShort') for s in submodels_of(env, shell_id)):
        raise ValidationError(f"The AAS already has a submodel named {source.get('idShort')}.")
    submodel = copy.deepcopy(source)
    submodel['id'], submodel['kind'] = new_id(), 'Instance'
    _replace_reference_targets(submodel, source['id'], submodel['id'])
    # Template qualifiers stay in the working copy for cardinality checks; finalize() removes them before export.
    if not keep_values:
        clear_values(submodel)
    _attach(env, shell_id, prune(submodel))
    known = {c['id'] for c in env.get('conceptDescriptions', [])}
    for description in template_env.get('conceptDescriptions', []):
        if description['id'] not in known:
            env.setdefault('conceptDescriptions', []).append(copy.deepcopy(description))
    return submodel


def add_element(env, shell_id, parent_target, model_type, id_short='', value_type='', semantic_id=''):
    parent = resolve(env, shell_id, parent_target)
    kind = parent.get('modelType')
    if kind not in CONTAINERS:
        raise ValidationError('Elements can only be added to submodels, collections, lists and entities.')
    if model_type not in NEW_ELEMENTS:
        raise ValidationError('Unsupported element type.')
    element = {'modelType': model_type}
    if kind == 'SubmodelElementList':
        if parent.get('typeValueListElement') != model_type:
            raise ValidationError(f"This list only holds {parent.get('typeValueListElement')} elements.")
        if model_type == 'Property':
            value_type = parent.get('valueTypeListElement') or value_type
    else:
        element['idShort'] = check_id_short(id_short)
        if any(c.get('idShort') == element['idShort'] for c in items_of(parent)):
            raise ValidationError(f"{element['idShort']} already exists here.")
    if model_type == 'Property':
        if value_type not in VALUE_TYPES:
            raise ValidationError('Select a value type for the property.')
        element['valueType'] = value_type
    if model_type == 'SubmodelElementList':
        element.update(typeValueListElement='SubmodelElementCollection', orderRelevant=True)
    if text(semantic_id):
        element['semanticId'] = external_reference(semantic_id)
    parent.setdefault(CHILDREN[kind], []).append(element)
    return element


def delete_element(env, shell_id, target):
    if not target.get('path'):
        submodel = resolve(env, shell_id, target)
        shell = get_shell(env, shell_id)
        shell['submodels'] = [r for r in shell.get('submodels', []) if r['keys'][0]['value'] != submodel['id']]
        if not any(r['keys'][0]['value'] == submodel['id'] for s in env.get('assetAdministrationShells', []) for r in s.get('submodels', [])):
            env['submodels'] = [s for s in env['submodels'] if s['id'] != submodel['id']]
        return
    parent = resolve(env, shell_id, dict(target, path=target['path'][:-1]))
    node = resolve(env, shell_id, target)
    items = items_of(parent)
    items.remove(node)


# ---------------------------------------------------------------- mapping rules
def parse_target(value):
    if not isinstance(value, dict) or not isinstance(value.get('path', []), list) or len(value.get('path', [])) > 64:
        raise ValidationError('A mapping target is invalid.')
    path = value.get('path', [])
    if any(not (isinstance(s, int) and not isinstance(s, bool) and s >= 0) and not (isinstance(s, str) and 0 < len(s) <= 128) for s in path):
        raise ValidationError('A mapping target path is invalid.')
    return {'submodel': text(value.get('submodel')), 'submodel_idShort': text(value.get('submodel_idShort')), 'path': path}


def normalize_rules(rules):
    if not isinstance(rules, list) or len(rules) > 500:
        raise ValidationError('Mappings must be a list of at most 500 rules.')
    result = []
    for rule in rules:
        if not isinstance(rule, dict) or rule.get('kind') not in ('value', 'rows'):
            raise ValidationError('Each mapping must be a single-value or table-row mapping.')
        sheet = text(rule.get('sheet'))
        if not sheet:
            raise ValidationError('Each mapping needs a worksheet.')
        item = {'kind': rule['kind'], 'sheet': sheet, 'target': parse_target(rule.get('target'))}
        if rule['kind'] == 'value':
            item['column'] = text(rule.get('column'))
            row, match = rule.get('row'), rule.get('match')
            if not item['column'] or (row is None) == (match is None):
                raise ValidationError('A single-value mapping needs a column and either a row number or a match condition.')
            if row is not None:
                if type(row) is not int or row < 2:
                    raise ValidationError('The row number must be an Excel row number of 2 or higher.')
                item['row'] = row
            else:
                if not isinstance(match, dict) or not text(match.get('column')) or not text(match.get('value')):
                    raise ValidationError('A match condition needs a column and a value.')
                item['match'] = {'column': text(match['column']), 'value': text(match['value'])}
            item['language'] = text(rule.get('language')) or 'en'
            if not LANGUAGE.fullmatch(item['language']):
                raise ValidationError('Invalid language tag.')
        else:
            item['key_column'] = text(rule.get('key_column'))
            item['mode'] = rule.get('mode', 'merge')
            if item['mode'] not in ('merge', 'replace'):
                raise ValidationError('The row import mode must be merge or replace.')
            prototype = rule.get('prototype')
            if prototype is not None and not (isinstance(prototype, int) and not isinstance(prototype, bool) or isinstance(prototype, str)):
                raise ValidationError('The row template element is invalid.')
            item['prototype'] = prototype
            columns = rule.get('columns')
            if not isinstance(columns, list) or not 1 <= len(columns) <= 256:
                raise ValidationError('A table-row mapping needs between 1 and 256 column mappings.')
            item['columns'] = []
            for column in columns:
                if not isinstance(column, dict) or not text(column.get('column')) or not isinstance(column.get('path'), list) or not 1 <= len(column['path']) <= 16:
                    raise ValidationError('Each column mapping needs a column and a target path.')
                path = column['path'] if column['path'] == [GLOBAL_ASSET] else [check_id_short(s) for s in column['path']]
                entry = {'column': text(column['column']), 'path': path}
                if column.get('valueType'):
                    if column['valueType'] not in VALUE_TYPES:
                        raise ValidationError('Unsupported value type.')
                    entry['valueType'] = column['valueType']
                if text(column.get('semanticId')):
                    external_reference(column['semanticId'])
                    entry['semanticId'] = text(column['semanticId'])
                item['columns'].append(entry)
        result.append(item)
    return result


def to_xsd(value, value_type):
    """Convert a cell to the lexical form of an XSD type; None means the cell is empty."""
    if isinstance(value, Uncached):
        raise ValidationError(f'Formula {value} has no saved result; open and save the file in Excel or paste values.')
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if value_type == 'xs:boolean':
        lowered = text(value).lower()
        if isinstance(value, bool) or lowered in ('true', 'false', '1', '0'):
            result = 'true' if value is True or lowered in ('true', '1') else 'false'
        else:
            raise ValidationError(f'Expected a boolean (true/false/1/0); got {value!r}')
    elif value_type in ('xs:double', 'xs:float'):
        result = repr(number(value))
    elif value_type == 'xs:decimal':
        if isinstance(value, bool) or not DECIMAL.fullmatch(text(value)):
            raise ValidationError(f'Expected a decimal number; got {value!r}')
        try:
            result = format(Decimal(repr(value) if isinstance(value, float) else text(value)), 'f')
        except InvalidOperation:
            raise ValidationError(f'Expected a decimal number; got {value!r}') from None
    elif value_type in ('xs:integer', 'xs:int', 'xs:long', 'xs:short', 'xs:byte', 'xs:nonNegativeInteger', 'xs:positiveInteger',
                        'xs:nonPositiveInteger', 'xs:negativeInteger', 'xs:unsignedInt', 'xs:unsignedLong', 'xs:unsignedShort', 'xs:unsignedByte'):
        result = str(integer(value))
    elif value_type == 'xs:date' and isinstance(value, (datetime.datetime, datetime.date)):
        result = (value.date() if isinstance(value, datetime.datetime) else value).isoformat()
    elif value_type == 'xs:dateTime' and isinstance(value, datetime.datetime):
        result = value.isoformat()
    elif value_type == 'xs:time' and isinstance(value, (datetime.time, datetime.datetime)):
        result = (value.time() if isinstance(value, datetime.datetime) else value).isoformat()
    elif isinstance(value, (datetime.date, datetime.time)):
        result = value.isoformat()
    else:
        result = text(value)
    if not verification.value_consistent_with_xsd_type(result, aas_types.DataTypeDefXSD(value_type)):
        raise ValidationError(f'{value!r} is not a valid {value_type} value.')
    return result


def _write(node, cell, language):
    """Write one cell into a Property, a MultiLanguageProperty or an Entity's global asset ID; returns False for empty cells."""
    if node.get('modelType') == 'Property':
        value = to_xsd(cell, node.get('valueType', 'xs:string'))
        if value is None:
            return False
        node['value'] = value
    elif node.get('modelType') == 'MultiLanguageProperty':
        value = to_xsd(cell, 'xs:string')
        if value is None:
            return False
        texts = [t for t in node.get('value', []) if t.get('language') != language]
        node['value'] = texts + [{'language': language, 'text': value}]
    elif node.get('modelType') == 'Entity':
        value = to_xsd(cell, 'xs:string')
        if value is None:
            return False
        node['globalAssetId'], node['entityType'] = value, 'SelfManagedEntity'
    else:
        raise ValidationError(f"Values can only be written to properties or entity asset IDs; the target is a {node.get('modelType')}.")
    return True


def _resolve_relative(node, path):
    if path == [GLOBAL_ASSET]:
        return node if node.get('modelType') == 'Entity' else None
    for segment in path:
        node = find_child(node, segment) if node is not None and node.get('modelType') != 'SubmodelElementList' else None
    return node


def apply_rules(env, shell_id, tables, rules):
    """Apply mapping rules to a copy of env. Returns the new environment and a report; errors leave nothing half-written."""
    env = copy.deepcopy(env)
    rules = normalize_rules(rules)
    report = {'errors': [], 'warnings': [], 'values': 0, 'created': 0, 'updated': 0, 'rules': []}
    def error(index, message, sheet='', row=None, column=''):
        report['errors'].append(dict(rule=index, message=message, sheet=sheet, row=row, column=column))
    for index, rule in enumerate(rules, 1):
        stats = {'rule': index, 'values': 0, 'created': 0, 'updated': 0, 'empty': 0}
        report['rules'].append(stats)
        sheet = tables.sheets.get(rule['sheet'])
        if sheet is None or sheet['problem']:
            error(index, f"Worksheet {rule['sheet']} is not available" + (f": {sheet['problem']}" if sheet else '.'), rule['sheet'])
            continue
        needed = [rule['column']] if rule['kind'] == 'value' else [c['column'] for c in rule['columns']] + ([rule['key_column']] if rule['key_column'] else [])
        if rule['kind'] == 'value' and 'match' in rule:
            needed.append(rule['match']['column'])
        missing = [c for c in needed if c not in sheet['columns']]
        if missing:
            error(index, 'Missing columns: ' + ', '.join(missing), rule['sheet'])
            continue
        try:
            target = resolve(env, shell_id, rule['target'])
        except ValidationError as exc:
            error(index, str(exc), rule['sheet'])
            continue
        if rule['kind'] == 'value':
            if 'row' in rule:
                rows = [r for r in sheet['rows'] if r['_row'] == rule['row']]
            else:
                rows = [r for r in sheet['rows'] if text(r.get(rule['match']['column'])) == rule['match']['value']]
            if len(rows) != 1:
                where = f"row {rule['row']}" if 'row' in rule else f"{rule['match']['column']} = {rule['match']['value']}"
                error(index, f'{len(rows)} rows match {where}; exactly one is required.', rule['sheet'])
                continue
            row = rows[0]
            try:
                if _write(target, row.get(rule['column']), rule['language']):
                    stats['values'] += 1
                else:
                    stats['empty'] += 1
                    report['warnings'].append(dict(rule=index, message='The cell is empty; the target keeps its current value.',
                                                   sheet=rule['sheet'], row=row['_row'], column=rule['column']))
            except ValidationError as exc:
                error(index, str(exc), rule['sheet'], row['_row'], rule['column'])
            continue
        _apply_rows(env, rule, index, sheet, target, stats, error)
    for stats in report['rules']:
        for key in ('values', 'created', 'updated'):
            report[key] += stats[key]
        if stats['empty'] and rules[stats['rule'] - 1]['kind'] == 'rows':
            report['warnings'].append(dict(rule=stats['rule'], message=f"{stats['empty']} empty cells were left without a value.",
                                           sheet=rules[stats['rule'] - 1]['sheet'], row=None, column=''))
        if stats.pop('unchanged', False):
            report['warnings'].append(dict(rule=stats['rule'], message='The worksheet has no data rows; the target was left unchanged.',
                                           sheet=rules[stats['rule'] - 1]['sheet'], row=None, column=''))
    return prune(env), report


def _apply_rows(env, rule, index, sheet, container, stats, error):
    kind = container.get('modelType')
    if kind not in CONTAINERS:
        error(index, f'Rows can only be imported into submodels, collections, lists or entities; the target is a {kind}.', rule['sheet'])
        return
    is_list = kind == 'SubmodelElementList'
    if is_list and container.get('typeValueListElement') not in ROW_ITEMS:
        error(index, 'Rows can only be imported into lists of collections or entities.', rule['sheet'])
        return
    if not sheet['rows']:
        # An empty sheet must not wipe the container, which would also remove a template's row structure.
        stats['unchanged'] = True
        return
    prototype = None
    if rule['prototype'] is not None:
        prototype = find_child(container, rule['prototype'])
        if prototype is None or prototype.get('modelType') not in ROW_ITEMS:
            error(index, f"The row template element {rule['prototype']} is missing or is not a collection or entity.", rule['sheet'])
            return
        prototype = clear_values(copy.deepcopy(prototype))
    blank = {'modelType': container['typeValueListElement'] if is_list else 'SubmodelElementCollection'}
    if blank['modelType'] == 'Entity':
        blank['entityType'] = 'CoManagedEntity'
    items = [] if is_list or rule['mode'] == 'replace' else list(items_of(container))
    by_id = {c.get('idShort'): c for c in items}
    first_row = {}
    for row in sheet['rows']:
        if is_list:
            item = copy.deepcopy(prototype or blank)
            item.pop('idShort', None)
            items.append(item)
            stats['created'] += 1
        else:
            key = text(row.get(rule['key_column'])) if rule['key_column'] else f"Row_{row['_row']}"
            if not key:
                error(index, 'The row name column is empty.', rule['sheet'], row['_row'], rule['key_column'])
                continue
            id_short = safe_id(key)
            if id_short in first_row:
                error(index, f"Rows {first_row[id_short]} and {row['_row']} both map to the name {id_short}; row names must be unique.",
                      rule['sheet'], row['_row'], rule['key_column'])
                continue
            first_row[id_short] = row['_row']
            item = by_id.get(id_short)
            if item is not None and item.get('modelType') not in ROW_ITEMS:
                error(index, f'{id_short} already exists and is not a collection or entity.', rule['sheet'], row['_row'], rule['key_column'])
                continue
            if item is None:
                item = copy.deepcopy(prototype or blank)
                item['idShort'] = id_short
                items.append(item)
                by_id[id_short] = item
                stats['created'] += 1
            else:
                stats['updated'] += 1
        for column in rule['columns']:
            node = _resolve_relative(item, column['path'])
            if node is None:
                if len(column['path']) != 1 or not column.get('valueType') or column['path'] == [GLOBAL_ASSET]:
                    error(index, 'Target ' + '/'.join(column['path']) + ' does not exist in the row element.', rule['sheet'], row['_row'], column['column'])
                    continue
                node = {'modelType': 'Property', 'idShort': column['path'][0], 'valueType': column['valueType']}
                if column.get('semanticId'):
                    node['semanticId'] = external_reference(column['semanticId'])
                item.setdefault(CHILDREN[item['modelType']], []).append(node)
            try:
                if _write(node, row.get(column['column']), 'en'):
                    stats['values'] += 1
                else:
                    stats['empty'] += 1
            except ValidationError as exc:
                error(index, str(exc), rule['sheet'], row['_row'], column['column'])
    container[CHILDREN[kind]] = items


def compare_trees(before, after):
    """Mark nodes of the result tree that are new or whose value changed."""
    old = {json.dumps(n['target'], sort_keys=True): n['value'] for n in before}
    for node in after:
        key = json.dumps(node['target'], sort_keys=True)
        node['changed'] = key not in old or old[key] != node['value']
    return after
