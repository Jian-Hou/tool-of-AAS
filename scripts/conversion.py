"""One validated Excel-to-AAS pipeline, with explicit incomplete-data reporting."""
import copy
import hashlib
import io
import json
import math
import os
import re
import shutil
import tempfile
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import quote, urlparse
from xml.etree import ElementTree as ET

import openpyxl
from aas_core3 import jsonization, verification
from schema import FAMILIES, CATALOG_NAMES, FIELDS, FIELD_BY_PATH, DEFAULT_SETTINGS, default_settings, default_bindings

MAX_ZIP_BYTES = 64 * 1024 * 1024
MAX_ROWS = 25000
TD = 'https://admin-shell.io/ZVEI/TechnicalData/'
RESERVED = {'CON', 'PRN', 'AUX', 'NUL'} | {f'{p}{i}' for p in ('COM', 'LPT') for i in range(1, 10)}

class ValidationError(ValueError):
    def __init__(self, message, report=None):
        super().__init__(message)
        self.report = report

@dataclass
class WorkbookData:
    sheets: dict
    headers: dict
    source_name: str
    source_sha256: str
    source_bytes: bytes

def text(value):
    return '' if value is None else str(value).strip()

def number(value):
    try:
        if isinstance(value, bool) or not text(value):
            raise ValueError()
        result = float(value)
        if not math.isfinite(result):
            raise ValueError()
        return result
    except (ValueError, TypeError, OverflowError):
        raise ValidationError(f'Expected a finite number; got {value!r}') from None

def integer(value):
    try:
        if isinstance(value, bool) or not text(value):
            raise InvalidOperation()
        result = Decimal(str(value))
        if not result.is_finite() or result != result.to_integral_value():
            raise InvalidOperation()
        return int(result)
    except (InvalidOperation, ValueError, TypeError):
        raise ValidationError(f'Expected an integer; got {value!r}') from None

def parse_vector(value, count):
    raw = text(value)
    if raw.startswith('(') and raw.endswith(')'):
        raw = raw[1:-1]
    elif '(' in raw or ')' in raw:
        raise ValidationError('Mismatched parentheses in numeric tuple')
    parts = raw.split(',')
    if len(parts) != count:
        raise ValidationError(f'Expected {count} numeric values; got {len(parts)}')
    return [number(part) for part in parts]

def parse_placement(value):
    raw = text(value)
    if '|' in raw:
        parts = raw.split('|')
        if len(parts) != 2:
            raise ValidationError('placement must contain a quaternion | six bounding box values')
        return parse_vector(parts[0], 4) + parse_vector(parts[1], 6)
    return parse_vector(raw, 10)

def safe_id(value):
    raw = text(value)
    result = re.sub(r'[^A-Za-z0-9_]', '_', raw)
    if not result or not result[0].isalpha():
        result = 'n_' + result
    # Preserve original valid IDs; transformed IDs carry a collision-resistant suffix.
    if result != raw or len(result) > 100:
        result = result[:90] + '_' + hashlib.sha256(raw.encode()).hexdigest()[:12]
    return result

def param_id(value):
    return re.sub(r'[^A-Za-z0-9_]', '_', text(value))

def parse_params(value):
    result, ids = {}, set()
    for part in text(value).split(';'):
        if not part.strip():
            continue
        if '=' not in part:
            raise ValidationError(f'Expected parameter name=value: {part!r}')
        key, raw = [p.strip() for p in part.split('=', 1)]
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_. -]*', key) or key in result:
            raise ValidationError(f'Invalid or duplicate parameter name: {key!r}')
        vals = [number(v) for v in raw.split(',')]
        expanded = [param_id(key)] if len(vals) == 1 else [f'{param_id(key)}_{i+1}' for i in range(len(vals))]
        if any(k in ids for k in expanded):
            raise ValidationError(f'Parameter names collide after normalization: {key}')
        ids.update(expanded)
        result[key] = vals
    return result

def read_workbook(path, source_name=None):
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValidationError('Only .xlsx is supported. Save legacy .xls files as .xlsx first.')
    payload = path.read_bytes()
    if len(payload) > 12 * 1024 * 1024:
        raise ValidationError('The Excel file exceeds the 12 MB limit.')
    try:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_ZIP_BYTES or len(z.infolist()) > 1000:
                raise ValidationError('The Excel archive exceeds the supported limits (64 MB / 1,000 entries).')
            if z.testzip():
                raise ValidationError('The Excel file is corrupt.')
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Unable to read Excel. Check that the file is a valid .xlsx workbook.') from exc
    sheets, headers, iterator = {}, {}, None
    try:
        if len(wb.sheetnames) > 32:
            raise ValidationError('The workbook contains more than 32 worksheets.')
        for ws in wb:
            if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > 128:
                raise ValidationError(f'{ws.title} exceeds the supported limits (25,000 rows / 128 columns).')
            iterator = ws.iter_rows()
            first = next(iterator, ())
            names = [text(c.value) for c in first]
            # Ignore truly empty trailing columns, but never unnamed populated columns.
            while names and not names[-1]:
                names.pop()
            if not names:
                raise ValidationError(f'{ws.title} has no header row.')
            if any(not n for n in names) or len(set(names)) != len(names):
                raise ValidationError(f'{ws.title} contains blank or duplicate column headers.')
            if any(n.startswith('_') for n in names):
                raise ValidationError(f'Column headers in {ws.title} must not start with an underscore.')
            rows = []
            for line, cells in enumerate(iterator, 2):
                if line > MAX_ROWS or len(cells) > 128:
                    raise ValidationError(f'{ws.title} exceeds the supported row or column limit.')
                if any(c.data_type == 'f' for c in cells):
                    raise ValidationError(f'{ws.title}, row {line}: formulas are not supported. Paste the conversion data as values.')
                if all(c.value in ('', None) for c in cells):
                    continue
                if any(c.value not in ('', None) for c in cells[len(names):]):
                    raise ValidationError(f'{ws.title}, row {line}: data exists in a column without a header.')
                row = {name: cells[i].value if i < len(cells) and cells[i].value is not None else '' for i, name in enumerate(names)}
                row['_row'] = line
                rows.append(row)
            sheets[ws.title], headers[ws.title] = rows, names
    finally:
        if iterator is not None and hasattr(iterator,'close'):
            iterator.close()
        wb.close()
    if 'components' not in sheets or 'joint_instances' not in sheets:
        raise ValidationError('The components and joint_instances worksheets are required. Sheet names are fixed; column names can be mapped.')
    return WorkbookData(sheets, headers, source_name or path.name, hashlib.sha256(payload).hexdigest(), payload)

def normalize_settings(settings=None):
    result = default_settings()
    settings = {} if settings is None else settings
    if not isinstance(settings, dict) or set(settings) - set(result):
        raise ValidationError('Invalid settings format or unknown settings fields.')
    for key, value in settings.items():
        if key in ('coordinate_system', 'parameter_units'):
            if not isinstance(value, dict):
                raise ValidationError(f'{key} must be an object.')
            if key == 'coordinate_system' and set(value) - set(result[key]):
                raise ValidationError('Coordinate settings contain unknown fields.')
            result[key].update(value)
        else:
            if not isinstance(value, str) or len(value) > 2000:
                raise ValidationError(f'{key} must be a string of at most 2,000 characters.')
            result[key] = value.strip()
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,63}', result['name']) or result['name'].upper() in RESERVED:
        raise ValidationError('The assembly name must start with a letter, contain only ASCII letters, digits and underscores, and be at most 64 characters long.')
    if result['asset_id'] and (not urlparse(result['asset_id']).scheme or re.search(r'\s', result['asset_id'])):
        raise ValidationError('The global asset ID must be a URI without whitespace, such as urn:... or https://....')
    cs = result['coordinate_system']
    if type(cs['Confirmed']) is not bool:
        raise ValidationError('Coordinate confirmation must be a boolean.')
    for k, v in cs.items():
        if k != 'Confirmed' and (not isinstance(v, str) or len(v) > 500):
            raise ValidationError(f'Invalid format for coordinate setting {k}.')
    for k, allowed in {
        'LengthUnit': ('', 'mm', 'cm', 'm', 'um'),
        'Handedness': ('', 'RightHanded', 'LeftHanded'),
        'QuaternionOrder': ('XYZW', 'WXYZ'),
        'RotationDirection': ('', 'LocalToAssembly', 'AssemblyToLocal'),
        'BoundingBoxCoordinateSystem': ('', 'ComponentLocal', 'AssemblyGlobal'),
    }.items():
        if cs[k] not in allowed:
            raise ValidationError(f'Unsupported value for coordinate setting {k}.')
    if cs['Confirmed'] and any(not cs[k].strip() for k in cs if k != 'Confirmed'):
        raise ValidationError('Complete all coordinate settings and their source before confirming the conventions.')
    if len(result['parameter_units']) > 500 or any(not isinstance(k, str) or not isinstance(v, str) or len(k) > 100 or len(v) > 40 for k,v in result['parameter_units'].items()):
        raise ValidationError('Invalid parameter unit settings.')
    return result

def normalize_bindings(bindings=None):
    if bindings is None or bindings == []:
        return default_bindings()
    if not isinstance(bindings, list) or len(bindings) > len(FIELDS):
        raise ValidationError('Field mappings must be a valid list.')
    by_path = {b['aas_path']: b for b in default_bindings()}
    seen = set()
    for b in bindings:
        if not isinstance(b, dict) or b.get('aas_path') not in FIELD_BY_PATH:
            raise ValidationError('A mapping contains an unknown target. Select a target from the current mapping contract.')
        target = b['aas_path']
        if target in seen:
            raise ValidationError('A target field cannot be mapped twice.')
        seen.add(target)
        if b.get('excel_sheet') != FIELD_BY_PATH[target]['excel_sheet']:
            raise ValidationError('Mappings must stay within the same worksheet. Records across sheets are joined by component label.')
        column = b.get('excel_column')
        if not isinstance(column, str) or not column.strip() or len(column) > 200:
            raise ValidationError('The source column name must be nonempty and at most 200 characters long.')
        if b.get('transform'):
            raise ValidationError('Custom transformation expressions are not supported.')
        by_path[target] = {k:b[k] for k in ('aas_path', 'excel_sheet', 'excel_column')}
    return list(by_path.values())

def analyze(workbook, settings=None, bindings=None):
    settings = normalize_settings(settings)
    bindings = normalize_bindings(bindings)
    bmap = {b['aas_path']:b for b in bindings}
    issues, mapped = [], {}
    def issue(severity, code, message, sheet='', row=None, field='', component=''):
        item = dict(severity=severity, code=code, message=message, sheet=sheet, row=row, field=field, component=component)
        issues.append(item)
    for sheet, rawrows in workbook.sheets.items():
        specs = [f for f in FIELDS if f['excel_sheet'] == sheet]
        if not specs:
            continue
        for spec in specs:
            source = bmap[spec['aas_path']]['excel_column']
            if spec['required'] and source not in workbook.headers[sheet]:
                issue('error','missing_column', f'Missing source column {source} (target: {spec["label"]})',sheet,1,source)
        mapped[sheet] = [{'_row': r['_row'], **{s['excel_column']:r.get(bmap[s['aas_path']]['excel_column'],'') for s in specs}} for r in rawrows]
    comps, joints = mapped.get('components',[]), mapped.get('joint_instances',[])
    if not comps:
        issue('error','empty_components','The components sheet contains no valid components.','components')
    if not joints:
        issue('warning','empty_joints','No joint records were provided; the output will contain no component connections.','joint_instances')
    # Duplicates are diagnosed before dictionaries can hide them.
    for sheet, cols in [('components',('assembly_id','label')),('joint_instances',('joint_id',))] + [(f+'_instances',('assembly_id','label')) for f in FAMILIES] + [(f+'_types',('shape_type',)) for f in FAMILIES] + [('joint_types',('joint_type',))]:
        for col in cols:
            seen = set()
            for row in mapped.get(sheet,[]):
                key = text(row.get(col))
                if key and col in ('assembly_id','joint_id'):
                    try:key=str(integer(row.get(col)))
                    except ValidationError:pass
                if not key:
                    issue('error','missing_key',f'{col} must not be empty.',sheet,row['_row'],col)
                elif key in seen:
                    issue('error','duplicate_key',f'Duplicate {col}: {key}',sheet,row['_row'],col)
                seen.add(key)
    comp_by_label = {text(r.get('label')):r for r in comps}
    parsed, placements, coords, missing, catalog = {}, {}, {}, {}, {}
    for family in FAMILIES:
        sn = family + '_instances'
        catalog[family] = {text(r.get('shape_type')):r for r in mapped.get(family+'_types',[])}
        for row in mapped.get(sn,[]):
            label = text(row.get('label'))
            parent = comp_by_label.get(label)
            if parent is None or text(parent.get('type')) != family:
                issue('error','orphan_instance','The instance has no matching component in the same family.',sn,row['_row'],'label',label)
            else:
                try:
                    if integer(parent.get('assembly_id')) != integer(row.get('assembly_id')):
                        issue('error','identity_mismatch','The instance and component have different assembly_id values.',sn,row['_row'],'assembly_id',label)
                except ValidationError as exc:
                    issue('error','invalid_instance_id',str(exc),sn,row['_row'],'assembly_id',label)
            try:
                params = parse_params(row.get('params'))
                if not params:
                    raise ValidationError('Geometry parameters must not be empty.')
                parsed[label] = params
            except ValidationError as exc:
                issue('error','invalid_params',str(exc),sn,row['_row'],'params',label)
                params = {}
            typ = text(row.get('shape_type'))
            if not typ:
                issue('error','missing_geometry_type','The instance geometry type must not be empty.',sn,row['_row'],'shape_type',label)
            definition = catalog[family].get(typ)
            if definition:
                required = {k.strip() for k in text(definition.get('params_needed')).split(';') if k.strip()}
                if not required:
                    issue('warning','empty_parameter_definition',f'{typ} has no required parameter definition.',family+'_types',definition['_row'],'params_needed')
                absent = sorted(required-set(params))
                if absent:
                    missing[label] = absent
                    issue('warning','missing_parameters','Missing required parameters: '+', '.join(absent),sn,row['_row'],'params',label)
            else:
                # Report once per family/type; build a clearly marked observed catalog entry later.
                key = (family,typ)
                if not any(x.get('code')=='observed_type' and x['field']==f'{family}/{typ}' for x in issues):
                    issue('warning','observed_type',f'{family}/{typ} has no formal type definition. The catalog only records parameters observed in instances.',sn,None,f'{family}/{typ}')
        actual_counts = Counter(text(r.get('shape_type')) for r in mapped.get(sn,[]))
        for typ, row in catalog[family].items():
            if text(row.get('count')):
                try:
                    if integer(row['count']) != actual_counts[typ]:
                        issue('warning','count_mismatch',f'The source count for {typ} does not match the number of instances.',family+'_types',row['_row'],'count')
                except ValidationError as exc:
                    issue('error','invalid_count',str(exc),family+'_types',row['_row'],'count')
    for row in comps:
        label, family = text(row.get('label')), text(row.get('type'))
        if family not in FAMILIES:
            issue('error','unknown_family',f'Unsupported component family: {family}', 'components',row['_row'],'type',label)
        if not text(row.get('shape_type')):
            issue('error','missing_source_type','The original component type code must not be empty.','components',row['_row'],'shape_type',label)
        if label not in parsed:
            issue('error','missing_instance','The component has no valid instance parameter record.','components',row['_row'],'label',label)
        for col, parser in [('assembly_id',integer),('coord',lambda v:parse_vector(v,3)),('placement',parse_placement)]:
            try:
                value = parser(row.get(col))
                if col == 'coord': coords[label] = value
                if col == 'placement':
                    placements[label] = value
                    if abs(math.sqrt(sum(q*q for q in value[:4]))-1)>1e-3:
                        issue('error','invalid_quaternion','Quaternion norm must be 1 within a tolerance of 0.001.','components',row['_row'],col,label)
                    if any(a>b for a,b in zip(value[4:7],value[7:10])):
                        issue('error','invalid_bbox','A bounding box minimum exceeds its corresponding maximum.','components',row['_row'],col,label)
            except ValidationError as exc:
                issue('error','invalid_'+col,str(exc),'components',row['_row'],col,label)
    joint_catalog = {text(r.get('joint_type')):r for r in mapped.get('joint_types',[])}
    for row in joints:
        try: integer(row.get('joint_id'))
        except ValidationError as exc: issue('error','invalid_joint_id',str(exc),'joint_instances',row['_row'],'joint_id')
        typ = text(row.get('joint_type'))
        if not typ:
            issue('error','missing_joint_type','The joint type must not be empty.','joint_instances',row['_row'],'joint_type')
        elif typ not in joint_catalog:
            issue('warning','unknown_joint_type',f'Joint type {typ} has no catalog definition.','joint_instances',row['_row'],'joint_type')
        for side in ('side1','side2'):
            target = text(row.get(side+'_id'))
            if not target:
                if not (typ == 'GroundedJoint' and side == 'side2'):
                    issue('error','missing_joint_endpoint',f'{side} has no component reference.','joint_instances',row['_row'],side+'_id')
            elif target not in comp_by_label and target != 'Assembly001' and target != settings['name']:
                issue('error','unresolved_joint',f'Joint component {target} does not exist.','joint_instances',row['_row'],side+'_id')
            if typ != 'GroundedJoint' and not text(row.get(side+'_sub')):
                issue('error','missing_joint_feature',f'{side} has no joint feature.','joint_instances',row['_row'],side+'_sub')
    for key, label in [('manufacturer_name','Manufacturer name'),('product_designation','Manufacturer product designation'),('article_number','Manufacturer article number'),('order_code','Manufacturer order code')]:
        if not settings[key]:
            issue('warning','manufacturer_missing',f'{label} was not provided. The field is retained without a value; template information is incomplete.',field=key)
    if not settings['coordinate_system']['Confirmed']:
        issue('warning','unconfirmed_coordinates','Coordinate, unit and rotation conventions are not confirmed by the data source. Original coord/placement values are preserved.')
    param_names = sorted({key for p in parsed.values() for key in p})
    unknown_units = [key for key in param_names if not parameter_unit(key,settings)]
    if unknown_units:
        issue('warning','unknown_parameter_units','Units were not provided for these parameters: '+', '.join(unknown_units))
    report = {
        'status': 'invalid' if any(i['severity']=='error' for i in issues) else ('draft' if issues else 'complete'),
        'errors': [i for i in issues if i['severity']=='error'], 'warnings': [i for i in issues if i['severity']=='warning'],
        'counts': {'components':len(comps),'joints':len(joints),'parameters':sum(len(v) for p in parsed.values() for v in p.values()),
                   'components_missing_parameters':len(missing),'missing_parameter_fields':sum(map(len,missing.values()))},
        'parameter_names':param_names,
    }
    context = dict(sheets=mapped, params=parsed, placements=placements, coords=coords, missing=missing,catalog=catalog)
    return report, context, settings, bindings

def parameter_unit(key, settings):
    # A suffix explicitly present in the source takes precedence over user defaults.
    if key.endswith('_mm'):
        return 'mm'
    return text(settings['parameter_units'].get(key))

def prop(name,value=None,kind='xs:string',sem=None):
    obj = {'idShort':name,'modelType':'Property','valueType':kind}
    if value is not None and value != '':
        obj['value'] = str(integer(value)) if kind=='xs:integer' else str(number(value)) if kind=='xs:double' else str(value)
    if sem: obj['semanticId'] = ext(sem)
    return obj

def coll(name, children, sem=None):
    obj = {'idShort':name,'modelType':'SubmodelElementCollection'}
    if children: obj['value'] = children
    if sem: obj['semanticId'] = ext(sem)
    return obj

def ext(value):
    return {'type':'ExternalReference','keys':[{'type':'GlobalReference','value':value}]}

def model_ref(keys):
    return {'type':'ModelReference','keys':[{'type':typ,'value':value} for typ,value in keys]}

def reference(name, keys):
    return {'idShort':name,'modelType':'ReferenceElement','value':model_ref(keys)}

def build_model(workbook, settings=None, bindings=None, strict=False):
    report, ctx, settings, bindings = analyze(workbook,settings,bindings)
    if report['errors'] or (strict and report['warnings']):
        raise ValidationError('The input did not pass strict validation.' if strict else 'The input contains errors and cannot be exported.',report)
    name = settings['name']
    stem = 'urn:aas:assembly:' + str(uuid.uuid5(uuid.NAMESPACE_URL,workbook.source_sha256+':'+name))
    adid, tdid, aasid = stem+':AssemblyDefinition',stem+':TechnicalData',stem+':aas'
    adkey = [('Submodel',adid)]
    sheets, params = ctx['sheets'],ctx['params']
    concepts = {}
    def param_properties(label):
        children=[]
        for key, vals in params[label].items():
            unit = parameter_unit(key,settings)
            concept = stem + ':parameter:' + quote(key,safe='')
            cd = {'id':concept,'idShort':safe_id(key),'modelType':'ConceptDescription',
                  'description':[{'language':'en','text':f'Source parameter: {key}. '+('Unit provided: '+unit if unit else 'Unit not supplied; interpretation is incomplete.')}]}
            if unit:
                cd['embeddedDataSpecifications']=[{
                    'dataSpecification':ext('https://admin-shell.io/DataSpecificationTemplates/DataSpecificationIEC61360/3/0'),
                    'dataSpecificationContent':{'modelType':'DataSpecificationIec61360','preferredName':[{'language':'en','text':key}],
                                              'definition':[{'language':'en','text':'Geometry parameter from source workbook: '+key}],
                                              'dataType':'REAL_MEASURE','unit':unit},
                }]
            concepts[concept]=cd
            for i,value in enumerate(vals):
                pid = param_id(key) if len(vals)==1 else f'{param_id(key)}_{i+1}'
                children.append(prop(pid,value,'xs:double',concept))
        return children
    general=[]
    for key, sid, sem in [('manufacturer_name','ManufacturerName','0173-1#02-AAO677#002'),
                           ('article_number','ManufacturerArticleNumber','0173-1#02-AAO676#003'),
                           ('order_code','ManufacturerOrderCode','0173-1#02-AAO227#002')]:
        general.append(prop(sid,settings[key],sem=sem))
    designation={'idShort':'ManufacturerProductDesignation','modelType':'MultiLanguageProperty','semanticId':ext('0173-1#02-AAW338#001')}
    if settings['product_designation']: designation['value']=[{'language':'en','text':settings['product_designation']}]
    general.append(designation)
    class_items=[]
    for family,code in sorted({(text(r['type']),text(r['shape_type'])) for r in sheets['components']}):
        class_items.append(coll(safe_id(f'{family}_{code}'),[
            prop('ProductClassificationSystem','Project component types',sem=TD+'ProductClassificationSystem/1/1'),
            prop('ProductClassId',f'{family}/{code}',sem=TD+'ProductClassId/1/1')],TD+'ProductClassificationItem/1/1'))
    sections=[coll(safe_id(text(r['label'])),param_properties(text(r['label'])),TD+'MainSection/1/1') for r in sheets['components']]
    td_elements=[coll('GeneralInformation',general,TD+'GeneralInformation/1/1'),
                 coll('ProductClassifications',class_items,TD+'ProductClassifications/1/1'),
                 coll('TechnicalProperties',sections,TD+'TechnicalProperties/1/1')]
    type_groups=[]
    instance_by_label={text(r['label']):r for family in FAMILIES for r in sheets.get(family+'_instances',[])}
    for family in FAMILIES:
        inst=sheets.get(family+'_instances',[])
        catalog=ctx['catalog'][family]
        codes=sorted(set(catalog)|{text(r['shape_type']) for r in inst})
        items=[]
        for code in codes:
            row=catalog.get(code)
            observed=sorted({k for r in inst if text(r['shape_type'])==code for k in params[text(r['label'])]})
            entries=[prop('TypeCode',code),prop('InstanceCount',sum(text(r['shape_type'])==code for r in inst),'xs:integer'),
                     prop('DefinitionStatus','SourceCatalog' if row else 'ObservedOnly')]
            if row:
                entries += [prop('GeometricDescription',row.get('geometric_description')),prop('RequiredParameters',row.get('params_needed'))]
                if text(row.get('count')): entries.append(prop('SourceCount',row['count'],'xs:integer'))
            if observed: entries.append(prop('ObservedParameters','; '.join(observed)))
            items.append(coll(safe_id(code),entries))
        if items: type_groups.append(coll(CATALOG_NAMES[family],items))
    joint_types=[]
    for row in sheets.get('joint_types',[]):
        entries=[prop('TypeCode',row['joint_type']),prop('GeometricDescription',row.get('description'))]
        if text(row.get('count')): entries.append(prop('SourceCount',row['count'],'xs:integer'))
        joint_types.append(coll(safe_id(row['joint_type']),entries))
    if joint_types:type_groups.append(coll('JointTypes',joint_types))
    comps=[]
    for row in sheets['components']:
        label,family=text(row['label']),text(row['type'])
        geometry=text(instance_by_label[label]['shape_type'])
        identity=[prop('ComponentId',row['assembly_id'],'xs:integer'),prop('Label',label),prop('AssetTag',row.get('tag')),
                  prop('ComponentType',family),prop('ShapeTypeCode',row['shape_type']),prop('GeometryTypeCode',geometry)]
        vals=ctx['placements'][label]
        q=vals[:4] if settings['coordinate_system']['QuaternionOrder']=='XYZW' else vals[1:4]+vals[:1]
        status='MissingRequiredParameters' if label in ctx['missing'] else 'CompleteAgainstCatalog' if geometry in ctx['catalog'][family] else 'CatalogIncomplete'
        children=[coll('Identity',identity),coll('Position',[prop(k,v,'xs:double') for k,v in zip(('X','Y','Z'),ctx['coords'][label])]),
                  coll('Orientation',[prop(k,v,'xs:double') for k,v in zip(('Qx','Qy','Qz','Qw'),q)]),
                  coll('BoundingBox',[prop(k,v,'xs:double') for k,v in zip(('XMin','YMin','ZMin','XMax','YMax','ZMax'),vals[4:])]),
                  prop('SourceCoord',row['coord']),prop('SourcePlacement',row['placement']),
                  prop('SourceParameters',instance_by_label[label]['params']),
                  reference('TypeDefinition',adkey+[('SubmodelElementCollection','TypeCatalog'),('SubmodelElementCollection',CATALOG_NAMES[family]),('SubmodelElementCollection',safe_id(geometry))]),
                  reference('TechnicalParameters',[('Submodel',tdid),('SubmodelElementCollection','TechnicalProperties'),('SubmodelElementCollection',safe_id(label))]),
                  coll('DataQuality',[prop('ParameterCompleteness',status),prop('MissingParameters','; '.join(ctx['missing'].get(label,[])))])]
        comps.append(coll(safe_id(label),children))
    joints=[]
    for row in sheets.get('joint_instances',[]):
        children=[prop('JointId',row['joint_id'],'xs:integer'),prop('JointType',row['joint_type'])]
        for side,prefix in [('side1','Side1'),('side2','Side2')]:
            target=text(row[side+'_id'])
            if target:
                children.append(prop(prefix+'Component',target))
                keys=[('AssetAdministrationShell',aasid)] if target in ('Assembly001',name) and target not in instance_by_label else adkey+[('SubmodelElementCollection','Components'),('SubmodelElementCollection',safe_id(target))]
                children.append(reference(prefix+'Reference',keys))
            if text(row[side+'_sub']):children.append(prop(prefix+'Feature',row[side+'_sub']))
        joints.append(coll('Joint_'+str(integer(row['joint_id'])),children))
    cs=settings['coordinate_system']
    cs_props=[prop(k,v) for k,v in cs.items() if k!='Confirmed']
    cs_props += [prop('ConventionStatus','ConfirmedByUser' if cs['Confirmed'] else 'Unconfirmed'),
                 prop('BoundingBoxConvention','XMin,YMin,ZMin,XMax,YMax,ZMax'),prop('OrientationOutputOrder','XYZW')]
    quality=[prop('Status',report['status']),prop('MissingParameterComponents',report['counts']['components_missing_parameters'],'xs:integer'),
             prop('MissingParameterFields',report['counts']['missing_parameter_fields'],'xs:integer')]
    for i,issue in enumerate(report['warnings'],1):
        quality.append(coll(f'Issue_{i}',[prop('Code',issue['code']),prop('Message',issue['message']),prop('Component',issue['component']),
                                         prop('Sheet',issue['sheet']),prop('Row',issue['row'],'xs:integer')]))
    ad_elements=[coll('CoordinateSystem',cs_props),coll('TypeCatalog',type_groups),coll('Components',comps),coll('Joints',joints),
                 coll('DataQuality',quality),coll('SourceData',[prop('Filename',workbook.source_name),prop('SHA256',workbook.source_sha256),
                 {'idShort':'OriginalWorkbook','modelType':'File','contentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','value':'/aasx/files/source.xlsx'}])]
    model={'assetAdministrationShells':[{'id':aasid,'idShort':name,'modelType':'AssetAdministrationShell',
           'assetInformation':{'assetKind':'Instance','globalAssetId':settings['asset_id'] or stem+':asset'},
           'submodels':[model_ref([('Submodel',tdid)]),model_ref([('Submodel',adid)])]}],
           'submodels':[{'id':tdid,'idShort':'TechnicalData','modelType':'Submodel','kind':'Instance','semanticId':ext(TD+'Submodel/1/2'),'submodelElements':td_elements},
                        {'id':adid,'idShort':'AssemblyDefinition','modelType':'Submodel','kind':'Instance','semanticId':ext('urn:aas:assembly-definition:2'),'submodelElements':ad_elements}],
           'conceptDescriptions':list(concepts.values())}
    validate_model(model)
    return model, report

def validate_model(model):
    try:
        env=jsonization.environment_from_jsonable(model)
        failures=[f'{e.path}: {e.cause}' for e in verification.verify(env)]
    except Exception as exc:
        raise ValidationError('Unable to parse the AAS structure: '+str(exc)) from exc
    if failures:
        raise ValidationError('AAS metamodel validation failed: '+'; '.join(failures[:10]))
    # Check referential integrity independently from the SDK's structural rules.
    objects={o['id']:o for group in ('assetAdministrationShells','submodels','conceptDescriptions') for o in model.get(group,[])}
    def walk(node):
        if isinstance(node,dict):
            if node.get('type')=='ModelReference':
                keys=node['keys'];target=objects.get(keys[0]['value'])
                for key in keys[1:]:
                    candidates=target.get('submodelElements',target.get('value',[])) if target else []
                    target=next((c for c in candidates if c.get('idShort')==key['value']),None)
                if target is None: raise ValidationError('Unresolved internal AAS reference: '+str(keys))
            for child in node.values():walk(child)
        elif isinstance(node,list):
            for child in node:walk(child)
    walk(model)

def atomic_json(path, data):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix='.tmp_',dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:
            json.dump(data,f,ensure_ascii=False,indent=2,allow_nan=False);f.flush();os.fsync(f.fileno())
        os.replace(tmp,path)
    finally:
        Path(tmp).unlink(missing_ok=True)

def write_aasx(model, path, workbook, report, settings, bindings=None):
    validate_model(model)
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix='.aas_',suffix='.aasx',dir=path.parent);os.close(fd)
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    def rels(items):
        root=ET.Element('Relationships',xmlns=relns)
        for i,(target,typ) in enumerate(items,1):ET.SubElement(root,'Relationship',Id=f'R{i}',Type=typ,Target=target)
        return ET.tostring(root,encoding='utf-8',xml_declaration=True)
    types=ET.Element('Types',xmlns='http://schemas.openxmlformats.org/package/2006/content-types')
    for extn,content_type in [('rels','application/vnd.openxmlformats-package.relationships+xml'),('json','application/json'),('xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')]:
        ET.SubElement(types,'Default',Extension=extn,ContentType=content_type)
    ET.SubElement(types,'Override',PartName='/aasx/aasx-origin',ContentType='text/plain')
    try:
        with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
            z.writestr('aasx/aasx-origin','')
            z.writestr('aasx/data.json',json.dumps(model,ensure_ascii=False,indent=2,allow_nan=False))
            z.writestr('aasx/files/source.xlsx',workbook.source_bytes)
            z.writestr('aasx/files/validation.json',json.dumps(report,ensure_ascii=False,indent=2))
            z.writestr('aasx/files/config.json',json.dumps({'settings':normalize_settings(settings),'bindings':normalize_bindings(bindings)},ensure_ascii=False,indent=2))
            z.writestr('_rels/.rels',rels([('/aasx/aasx-origin','http://admin-shell.io/aasx/relationships/aasx-origin')]))
            z.writestr('aasx/_rels/aasx-origin.rels',rels([('/aasx/data.json','http://admin-shell.io/aasx/relationships/aas-spec')]))
            z.writestr('aasx/_rels/data.json.rels',rels([(f'/aasx/files/{n}','http://admin-shell.io/aasx/relationships/aas-suppl') for n in ('source.xlsx','validation.json','config.json')]))
            z.writestr('[Content_Types].xml',ET.tostring(types,encoding='utf-8',xml_declaration=True))
        with zipfile.ZipFile(tmp) as z:
            if z.testzip() or json.loads(z.read('aasx/data.json'))!=model:
                raise ValidationError('The exported file failed read-back verification.')
        if path.exists():
            backup=path.parent/'backups';backup.mkdir(exist_ok=True)
            shutil.copy2(path,backup/(path.stem+'_'+uuid.uuid4().hex+'.aasx'))
        os.replace(tmp,path)
    finally:
        Path(tmp).unlink(missing_ok=True)

def read_aasx(path):
    with zipfile.ZipFile(path) as z:
        if z.testzip():raise ValidationError('The AASX file is corrupt.')
        return json.loads(z.read('aasx/data.json'))
