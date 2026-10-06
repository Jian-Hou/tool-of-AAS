"""Workbook and export tests."""
import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

import openpyxl
from basyx.aas import model
from basyx.aas.adapter.aasx import AASXReader, DictSupplementaryFileContainer

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from conversion import ValidationError, analyze, build_model, read_aasx, read_workbook, write_aasx
from qa import audit
from schema import default_settings

# Minimal example workbook.
SHEETS = {
    'components': [['assembly_id', 'label', 'tag', 'type', 'shape_type', 'coord', 'placement'],
                   [1, 'demo_pipe_001', 'DEMO', 'pipeline', 'P01', '(0,0,0)', '(0,0,0,1)|0,0,0,100,10,10']],
    'pipeline_instances': [['assembly_id', 'label', 'shape_type', 'params'],
                           [1, 'demo_pipe_001', 'P01', 'length_mm=100;outer_diameter_mm=10;inner_diameter_mm=8']],
    'pipeline_types': [['shape_type', 'geometric_description', 'params_needed', 'count'],
                       ['P01', 'Synthetic demonstration pipe', 'length_mm;outer_diameter_mm;inner_diameter_mm', 1]],
    'joint_instances': [['joint_id', 'joint_type', 'side1_id', 'side1_sub', 'side2_id', 'side2_sub'],
                        [1, 'GroundedJoint', 'demo_pipe_001', None, None, None]],
    'joint_types': [['joint_type', 'description', 'count'], ['GroundedJoint', 'Synthetic grounding example', 1]],
}


def minimal_workbook(path, extra=None):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in SHEETS.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    if extra:
        extra(workbook)
    workbook.save(path)
    return Path(path)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(prefix='aas_workbook_')
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_unrelated_sheets_are_ignored(self):
        def extra(workbook):
            workbook.create_sheet('Sheet1')
            notes = workbook.create_sheet('notes')
            notes.append(['note', 'note', ''])
            notes.append(['total', '=1+1', 'unlabelled'])
        workbook = read_workbook(minimal_workbook(self.dir / 'extra.xlsx', extra))
        self.assertEqual(workbook.ignored_sheets, ('Sheet1', 'notes'))
        self.assertNotIn('notes', workbook.sheets)
        report = analyze(workbook)[0]
        self.assertEqual(report['status'], 'draft')
        self.assertEqual(report['ignored_sheets'], ['Sheet1', 'notes'])

    def test_recognized_sheets_are_still_validated(self):
        def formula(workbook):
            workbook['components'].append([2, 'demo_pipe_002', '=1+1', 'pipeline', 'P01', '(0,0,0)', '(0,0,0,1)|0,0,0,1,1,1'])
        with self.assertRaisesRegex(ValidationError, 'formulas'):
            read_workbook(minimal_workbook(self.dir / 'formula.xlsx', formula))
        with self.assertRaisesRegex(ValidationError, 'no header row'):
            read_workbook(minimal_workbook(self.dir / 'empty.xlsx', lambda workbook: workbook.create_sheet('tank_types')))


class ExportTests(unittest.TestCase):
    def test_exported_package_reads_back_with_basyx(self):
        with tempfile.TemporaryDirectory(prefix='aas_export_') as tmp:
            source = minimal_workbook(Path(tmp) / 'minimal.xlsx')
            workbook = read_workbook(source)
            aas_model, report = build_model(workbook)
            output = Path(tmp) / 'minimal.aasx'
            write_aasx(aas_model, output, workbook, report, default_settings())

            store, files = model.DictIdentifiableStore(), DictSupplementaryFileContainer()
            with AASXReader(output, failsafe=False) as reader:
                reader.read_into(store, files)
            self.assertEqual(len(store), 1 + 2 + 3)  # shell, 2 submodels, 3 concepts
            assembly = next(item for item in store if getattr(item, 'id_short', None) == 'AssemblyDefinition')
            self.assertEqual(assembly.get_referable(['Components', 'demo_pipe_001', 'Identity', 'Label']).value, 'demo_pipe_001')
            self.assertEqual(assembly.get_referable(['DataQuality', 'Status']).value, 'draft')
            embedded = io.BytesIO()
            files.write_file('/aasx/files/source.xlsx', embedded)
            self.assertEqual(embedded.getvalue(), source.read_bytes())
            with zipfile.ZipFile(output) as package:
                self.assertTrue({'aasx/files/validation.json', 'aasx/files/config.json'} <= set(package.namelist()))
            self.assertTrue(audit(workbook, read_aasx(output))['data_match'])


if __name__ == '__main__':
    unittest.main()
