"""Generic Excel-to-AAS mapping with synthetic workbooks, templates and packages built by the BaSyx SDK."""
import copy
import datetime
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

import openpyxl
from basyx.aas import model
from basyx.aas.adapter.aasx import AASXReader, AASXWriter, DictSupplementaryFileContainer

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import aas_mapper as m  # noqa: E402
from conversion import ValidationError, build_model, read_workbook, write_aasx  # noqa: E402
from schema import default_settings  # noqa: E402
from test_workbook_export import minimal_workbook  # noqa: E402

TEMPLATE_SEMANTIC = 'https://example.org/sm/Nameplate/1/0'


def workbook(path):
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = 'info'
    for row in [('key', 'value'), ('manufacturer', 'ACME GmbH'), ('serial', 'SN-001'), ('year', 2024)]:
        info.append(row)
    parts = wb.create_sheet('parts')
    parts.append(['part', 'material', 'mass_kg', 'count', 'inspected', 'delivered'])
    parts.append(['pipe-1', 'steel', 1.5, 2, True, datetime.datetime(2026, 1, 5)])
    parts.append(['valve_2', 'brass', '0.25', 1, False, datetime.datetime(2026, 2, 1)])
    parts.append(['tank 3', None, 12, 1, True, datetime.datetime(2026, 3, 9)])
    notes = wb.create_sheet('notes')
    notes.append(['note', 'total'])
    notes.append(['sum', '=1+1'])
    wb.create_sheet('empty')
    wb.save(path)
    return m.read_tables(path)


def template_package(path):
    cardinality = lambda value: {model.Qualifier('SMT/Cardinality', str, value=value, kind=model.QualifierKind.TEMPLATE_QUALIFIER)}
    marking = model.SubmodelElementCollection(None, value=[model.Property('MarkingName', str), model.Property('MarkingCount', model.datatypes.Int)])
    submodel = model.Submodel('urn:template:nameplate', id_short='Nameplate', kind=model.ModellingKind.TEMPLATE,
                              semantic_id=model.ExternalReference((model.Key(model.KeyTypes.GLOBAL_REFERENCE, TEMPLATE_SEMANTIC),)),
                              submodel_element=[
                                  model.MultiLanguageProperty('ManufacturerName', value=model.MultiLanguageTextType({'en': 'Example AG'}), qualifier=cardinality('One')),
                                  model.Property('SerialNumber', str, value='0000', qualifier=cardinality('ZeroToOne')),
                                  model.Property('YearOfConstruction', str),
                                  model.SubmodelElementList('Markings', model.SubmodelElementCollection, value=[marking], qualifier=cardinality('ZeroToMany')),
                              ])
    description = model.ConceptDescription('urn:cd:manufacturer-name', id_short='ManufacturerName')
    with AASXWriter(path) as writer:
        writer.write_all_aas_objects('/aasx/template.xml', model.DictIdentifiableStore([submodel, description]), DictSupplementaryFileContainer())
    return m.load_source(path)


def existing_package(path):
    files = DictSupplementaryFileContainer()
    import io
    name = files.add_file('/aasx/files/manual.pdf', io.BytesIO(b'%PDF-1.4 synthetic manual'), 'application/pdf')
    documentation = model.Submodel('urn:example:pump01:documentation', id_short='Documentation',
                                   submodel_element=[model.File('Manual', 'application/pdf', value=name), model.Property('Revision', str, value='B')])
    shell = model.AssetAdministrationShell(model.AssetInformation(model.AssetKind.INSTANCE, global_asset_id='urn:example:asset:pump01'),
                                           'urn:example:aas:pump01', id_short='Pump01', submodel={model.ModelReference.from_referable(documentation)})
    with AASXWriter(path) as writer:
        writer.write_aas(shell.id, model.DictIdentifiableStore([shell, documentation]), files, write_json=False)
        writer.write_thumbnail('/aasx/thumbnail.png', b'\x89PNG synthetic thumbnail', 'image/png')
    return m.load_source(path)


def read_back(path):
    store, files = model.DictIdentifiableStore(), DictSupplementaryFileContainer()
    with AASXReader(path, failsafe=False) as reader:
        reader.read_into(store, files)
    return store, files


def target(env, shell_id, submodel_id_short, *path):
    submodel = next(s for s in m.submodels_of(env, shell_id) if s['idShort'] == submodel_id_short)
    return {'submodel': submodel['id'], 'submodel_idShort': submodel_id_short, 'path': list(path)}


class MapperTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(prefix='aas_mapper_')
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        self.tables = workbook(self.dir / 'data.xlsx')

    def nameplate_env(self):
        env = m.new_environment('Pump01', 'urn:example:asset:pump01')
        shell_id = env['assetAdministrationShells'][0]['id']
        template = template_package(self.dir / 'template.aasx')
        m.add_template_submodel(env, shell_id, template.env, 'urn:template:nameplate')
        return env, shell_id

    def test_read_tables_reports_types_formulas_and_problems(self):
        sheets = m.table_summary(self.tables)
        self.assertEqual(sheets['parts']['types'], {'part': 'xs:string', 'material': 'xs:string', 'mass_kg': 'xs:string', 'count': 'xs:integer',
                                                    'inspected': 'xs:boolean', 'delivered': 'xs:date'})
        self.assertEqual(sheets['parts']['preview'][0]['_row'], 2)
        self.assertTrue(any('no saved result' in note for note in sheets['notes']['notes']))
        self.assertEqual(sheets['empty']['problem'], 'The first row contains no column headers.')

    def test_template_is_instantiated_without_example_values_or_template_qualifiers(self):
        env, shell_id = self.nameplate_env()
        nameplate = m.resolve(env, shell_id, target(env, shell_id, 'Nameplate'))
        self.assertEqual(nameplate['kind'], 'Instance')
        self.assertNotEqual(nameplate['id'], 'urn:template:nameplate')
        self.assertEqual(nameplate['semanticId']['keys'][0]['value'], TEMPLATE_SEMANTIC)
        self.assertNotIn('qualifiers', str(nameplate))
        self.assertNotIn('Example AG', str(nameplate))
        self.assertEqual([c['id'] for c in env['conceptDescriptions']], ['urn:cd:manufacturer-name'])
        self.assertEqual(m.verify(env), [])

    def test_new_aas_from_template_plus_custom_elements(self):
        env, shell_id = self.nameplate_env()
        m.add_element(env, shell_id, target(env, shell_id, 'Nameplate'), 'Property', 'Year', 'xs:int')
        m.add_submodel(env, shell_id, 'Parts')
        rules = [
            {'kind': 'value', 'sheet': 'info', 'column': 'value', 'match': {'column': 'key', 'value': 'manufacturer'},
             'target': target(env, shell_id, 'Nameplate', 'ManufacturerName'), 'language': 'de'},
            {'kind': 'value', 'sheet': 'info', 'column': 'value', 'row': 3, 'target': target(env, shell_id, 'Nameplate', 'SerialNumber')},
            {'kind': 'value', 'sheet': 'info', 'column': 'value', 'row': 4, 'target': target(env, shell_id, 'Nameplate', 'Year')},
            {'kind': 'rows', 'sheet': 'parts', 'key_column': 'part', 'target': target(env, shell_id, 'Parts'), 'columns': [
                {'column': 'part', 'path': ['Label'], 'valueType': 'xs:string'},
                {'column': 'material', 'path': ['Material'], 'valueType': 'xs:string'},
                {'column': 'mass_kg', 'path': ['MassKg'], 'valueType': 'xs:double', 'semanticId': 'https://example.org/mass'},
                {'column': 'inspected', 'path': ['Inspected'], 'valueType': 'xs:boolean'},
                {'column': 'delivered', 'path': ['Delivered'], 'valueType': 'xs:date'}]},
            {'kind': 'rows', 'sheet': 'parts', 'prototype': 0, 'target': target(env, shell_id, 'Nameplate', 'Markings'), 'columns': [
                {'column': 'part', 'path': ['MarkingName']}, {'column': 'count', 'path': ['MarkingCount']}]},
        ]
        result, report = m.apply_rules(env, shell_id, self.tables, rules)
        self.assertEqual(report['errors'], [], report)
        self.assertEqual((report['created'], report['values']), (6, 3 + 14 + 6))
        self.assertEqual(m.new_failures(env, result), [])
        out = self.dir / 'new.aasx'
        m.write_package(out, result)
        store, _ = read_back(out)
        nameplate = next(x for x in store if getattr(x, 'id_short', None) == 'Nameplate')
        self.assertEqual(nameplate.get_referable('ManufacturerName').value['de'], 'ACME GmbH')
        self.assertEqual(nameplate.get_referable('SerialNumber').value, 'SN-001')
        self.assertEqual(nameplate.get_referable('Year').value, 2024)
        markings = nameplate.get_referable('Markings').value
        self.assertEqual([(x.get_referable('MarkingName').value, x.get_referable('MarkingCount').value) for x in markings],
                         [('pipe-1', 2), ('valve_2', 1), ('tank 3', 1)])
        parts = next(x for x in store if getattr(x, 'id_short', None) == 'Parts')
        row = parts.get_referable(m.safe_id('pipe-1'))
        self.assertEqual((row.get_referable('Label').value, row.get_referable('MassKg').value, row.get_referable('Inspected').value),
                         ('pipe-1', 1.5, True))
        self.assertEqual(str(row.get_referable('Delivered').value), '2026-01-05')
        self.assertEqual(parts.get_referable('valve_2').get_referable('MassKg').value, 0.25)
        self.assertIsNone(parts.get_referable(m.safe_id('tank 3')).get_referable('Material').value)
        self.assertTrue(any('empty cells' in w['message'] for w in report['warnings']))

    def test_existing_xml_package_is_patched_in_place(self):
        source = self.dir / 'pump.aasx'
        package = existing_package(source)
        self.assertEqual(package.spec_format, 'xml')
        env, shell_id = package.env, 'urn:example:aas:pump01'
        m.add_submodel(env, shell_id, 'Parts')
        rules = [{'kind': 'rows', 'sheet': 'parts', 'key_column': 'part', 'target': target(env, shell_id, 'Parts'),
                  'columns': [{'column': 'count', 'path': ['Count'], 'valueType': 'xs:int'}]}]
        result, report = m.apply_rules(env, shell_id, self.tables, rules)
        self.assertEqual(report['errors'], [])
        out = self.dir / 'pump_out.aasx'
        m.write_package(out, result, base=source, package=package)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(out) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != package.spec_part:
                    self.assertEqual(before.read(name), after.read(name), name)
            self.assertTrue(after.read(package.spec_part).startswith(b'<?xml'))
        store, files = read_back(out)
        shell = store.get_item('urn:example:aas:pump01')
        self.assertEqual(shell.asset_information.global_asset_id, 'urn:example:asset:pump01')
        documentation = store.get_item('urn:example:pump01:documentation')
        self.assertEqual(documentation.get_referable('Revision').value, 'B')
        self.assertIn(documentation.get_referable('Manual').value, files)
        parts = next(x for x in store if getattr(x, 'id_short', None) == 'Parts')
        self.assertEqual(parts.get_referable(m.safe_id('pipe-1')).get_referable('Count').value, 2)

    def test_output_of_the_assembly_converter_can_be_extended(self):
        source = minimal_workbook(self.dir / 'assembly.xlsx')
        assembly = read_workbook(source)
        model_json, report = build_model(assembly)
        base = self.dir / 'assembly.aasx'
        write_aasx(model_json, base, assembly, report, default_settings())
        package = m.load_source(base)
        shell_id = package.env['assetAdministrationShells'][0]['id']
        rule = {'kind': 'value', 'sheet': 'info', 'column': 'value', 'row': 2,
                'target': target(package.env, shell_id, 'TechnicalData', 'GeneralInformation', 'ManufacturerName')}
        result, report = m.apply_rules(package.env, shell_id, self.tables, [rule])
        self.assertEqual(report['errors'], [])
        out = self.dir / 'assembly_out.aasx'
        m.write_package(out, result, base=base, package=package)
        with zipfile.ZipFile(base) as before, zipfile.ZipFile(out) as after:
            self.assertEqual(before.read('aasx/files/source.xlsx'), after.read('aasx/files/source.xlsx'))
        self.assertEqual(m.resolve(m.load_source(out).env, shell_id, rule['target'])['value'], 'ACME GmbH')

    def test_merge_updates_rows_and_replace_drops_stale_rows(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Parts')
        rule = {'kind': 'rows', 'sheet': 'parts', 'key_column': 'part', 'target': target(env, shell_id, 'Parts'),
                'columns': [{'column': 'count', 'path': ['Count'], 'valueType': 'xs:int'}]}
        first, _ = m.apply_rules(env, shell_id, self.tables, [rule])
        m.add_element(first, shell_id, target(first, shell_id, 'Parts'), 'Property', 'Manual', 'xs:string')
        again, report = m.apply_rules(first, shell_id, self.tables, [rule])
        parts = m.resolve(again, shell_id, target(again, shell_id, 'Parts'))
        self.assertEqual((report['created'], report['updated'], len(parts['submodelElements'])), (0, 3, 4))
        replaced, _ = m.apply_rules(first, shell_id, self.tables, [dict(rule, mode='replace')])
        self.assertEqual(len(m.resolve(replaced, shell_id, target(replaced, shell_id, 'Parts'))['submodelElements']), 3)

    def test_cell_errors_block_the_import_with_locations(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Data')
        m.add_element(env, shell_id, target(env, shell_id, 'Data'), 'Property', 'Count', 'xs:int')
        m.add_element(env, shell_id, target(env, shell_id, 'Data'), 'Property', 'Total', 'xs:string')
        rules = [
            {'kind': 'value', 'sheet': 'parts', 'column': 'part', 'row': 2, 'target': target(env, shell_id, 'Data', 'Count')},
            {'kind': 'value', 'sheet': 'notes', 'column': 'total', 'row': 2, 'target': target(env, shell_id, 'Data', 'Total')},
            {'kind': 'value', 'sheet': 'info', 'column': 'value', 'match': {'column': 'key', 'value': 'missing'}, 'target': target(env, shell_id, 'Data', 'Total')},
            {'kind': 'value', 'sheet': 'info', 'column': 'nope', 'row': 2, 'target': target(env, shell_id, 'Data', 'Total')},
            {'kind': 'value', 'sheet': 'info', 'column': 'value', 'row': 2, 'target': target(env, shell_id, 'Data', 'Missing')},
        ]
        _, report = m.apply_rules(env, shell_id, self.tables, rules)
        errors = {(e['rule'], e['row'], e['column']): e['message'] for e in report['errors']}
        self.assertIn('Expected an integer', errors[(1, 2, 'part')])
        self.assertIn('no saved result', errors[(2, 2, 'total')])
        self.assertIn('0 rows match', errors[(3, None, '')])
        self.assertIn('Missing columns: nope', errors[(4, None, '')])
        self.assertIn('Target not found', errors[(5, None, '')])

    def test_duplicate_row_names_and_empty_sheets(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Parts')
        wb = openpyxl.Workbook()
        wb.active.title = 'dup'
        for row in [('part', 'count'), ('pipe-1', 1), ('pipe-1 ', 2)]:
            wb.active.append(row)
        wb.create_sheet('none').append(['part', 'count'])
        wb.save(self.dir / 'edge.xlsx')
        tables = m.read_tables(self.dir / 'edge.xlsx')
        rule = {'kind': 'rows', 'sheet': 'dup', 'key_column': 'part', 'target': target(env, shell_id, 'Parts'),
                'columns': [{'column': 'count', 'path': ['Count'], 'valueType': 'xs:int'}]}
        _, report = m.apply_rules(env, shell_id, tables, [rule])
        self.assertIn('Rows 2 and 3 both map to the name', report['errors'][0]['message'])
        template = template_package(self.dir / 'template.aasx')
        m.add_template_submodel(env, shell_id, template.env, 'urn:template:nameplate')
        markings = target(env, shell_id, 'Nameplate', 'Markings')
        result, report = m.apply_rules(env, shell_id, tables, [dict(rule, sheet='none', target=markings, prototype=0, key_column='')])
        self.assertEqual(report['errors'], [])
        self.assertIn('left unchanged', report['warnings'][0]['message'])
        self.assertEqual(m.resolve(result, shell_id, markings), m.resolve(env, shell_id, markings))

    def test_large_containers_are_summarized_in_the_tree(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Parts')
        for i in range(30):
            m.add_element(env, shell_id, target(env, shell_id, 'Parts'), 'Property', f'P{i}', 'xs:string')
        nodes = m.tree(env, shell_id)
        self.assertEqual(len(nodes), 1 + 25 + 1)
        self.assertEqual((nodes[-1]['more'], nodes[-1]['label']), (5, '… 5 more'))

    def test_unreadable_packages_get_specific_messages(self):
        def package(spec_name, spec, target=None):
            path = self.dir / 'broken.aasx'
            rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r" Type="http://admin-shell.io/aasx/relationships/{}" Target="{}"/></Relationships>'
            with zipfile.ZipFile(path, 'w') as z:
                z.writestr('_rels/.rels', rels.format('aasx-origin', '/aasx/aasx-origin'))
                z.writestr('aasx/aasx-origin', '')
                z.writestr('aasx/_rels/aasx-origin.rels', rels.format('aas-spec', target or '/' + spec_name))
                z.writestr(spec_name, spec)
            return path
        with self.assertRaisesRegex(ValidationError, 'V1/V2 model'):
            m.load_source(package('aasx/data.xml', '<aasenv xmlns="http://www.admin-shell.io/aas/2/0"/>'))
        with self.assertRaisesRegex(ValidationError, 'is missing'):
            m.load_source(package('aasx/data.json', '{}', target='/aasx/other.json'))

    def test_existing_metamodel_problems_are_not_blamed_on_the_import(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Data')
        m.add_element(env, shell_id, target(env, shell_id, 'Data'), 'Property', 'Broken', 'xs:int')
        m.add_element(env, shell_id, target(env, shell_id, 'Data'), 'Property', 'Serial', 'xs:string')
        m.resolve(env, shell_id, target(env, shell_id, 'Data', 'Broken'))['value'] = 'not a number'
        self.assertTrue(m.verify(env))
        rule = {'kind': 'value', 'sheet': 'info', 'column': 'value', 'row': 3, 'target': target(env, shell_id, 'Data', 'Serial')}
        result, report = m.apply_rules(env, shell_id, self.tables, [rule])
        self.assertEqual((report['errors'], m.new_failures(env, result)), ([], []))

    def test_structure_edits_are_validated(self):
        env = m.new_environment('Line')
        shell_id = env['assetAdministrationShells'][0]['id']
        m.add_submodel(env, shell_id, 'Data')
        data = target(env, shell_id, 'Data')
        for args in [('Property', '1bad', 'xs:string'), ('Property', 'NoType', ''), ('Range', 'R', 'xs:int')]:
            with self.subTest(args=args), self.assertRaises(ValidationError):
                m.add_element(env, shell_id, data, *args)
        m.add_element(env, shell_id, data, 'SubmodelElementList', 'Items')
        with self.assertRaises(ValidationError):
            m.add_element(env, shell_id, data, 'Property', 'Items', 'xs:string')
        with self.assertRaises(ValidationError):
            m.add_element(env, shell_id, target(env, shell_id, 'Data', 'Items'), 'Property', '', 'xs:string')
        m.add_element(env, shell_id, target(env, shell_id, 'Data', 'Items'), 'SubmodelElementCollection')
        self.assertEqual(m.verify(env), [])
        m.delete_element(env, shell_id, target(env, shell_id, 'Data', 'Items', 0))
        m.delete_element(env, shell_id, data)
        self.assertEqual((m.submodels_of(env, shell_id), m.verify(env)), ([], []))
        with self.assertRaises(ValidationError):
            m.normalize_rules([{'kind': 'rows', 'sheet': 'parts', 'target': data, 'columns': [{'column': 'part', 'path': ['bad name']}]}])


if __name__ == '__main__':
    unittest.main()
